"""Viettel Phú Hà storefront, JSON API and protected administration."""
import argparse
import csv
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
import time
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from flask import Flask, Response, g, jsonify, redirect, request, send_file
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('DATA_DIR', str(ROOT / 'data'))).resolve()
DB = DATA / 'store.db'
BRANCHES = [
    {'id': 'cs1', 'name': 'CS1', 'address': 'KCN Phú Hà, thị xã Phú Thọ, Phú Thọ', 'map_url': ''},
    {'id': 'cs2', 'name': 'CS2', 'address': 'Thanh Miếu, Việt Trì, Phú Thọ', 'map_url': ''},
    {'id': 'cs3', 'name': 'CS3', 'address': 'Cao Mại, Lâm Thao, Phú Thọ', 'map_url': ''},
]
DEFAULT_SETTINGS = {
    'name': 'Viettel Phú Hà', 'phone': '0344.639.666', 'email': 'thuntlpto.9x@gmail.com',
    'hours': '', 'zalo_url': '', 'site_url': '', 'banner_url': '',
    'tagline': 'Kết nối giá trị - Đồng hành cùng bạn',
    'branches': BRANCHES, 'registration_policy': '', 'privacy_policy': '',
}
app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024
if os.environ.get('TRUST_PROXY') == '1':
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)


class Problem(Exception):
    def __init__(self, message, status=400):
        self.message, self.status = message, status


def now():
    return int(time.time())


@contextmanager
def database(write=False):
    db = sqlite3.connect(DB, timeout=15)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    try:
        if write:
            db.execute('BEGIN IMMEDIATE')
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def initialize():
    DATA.mkdir(parents=True, exist_ok=True)
    with database() as db:
        db.execute('PRAGMA journal_mode=WAL')
        db.executescript('''
        CREATE TABLE IF NOT EXISTS settings (id INTEGER PRIMARY KEY CHECK(id=1), value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
          password TEXT NOT NULL, role TEXT NOT NULL, branch TEXT, active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS sessions (token TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
          csrf TEXT NOT NULL, expires INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS sims (id INTEGER PRIMARY KEY, number TEXT UNIQUE NOT NULL, price INTEGER NOT NULL,
          type TEXT NOT NULL, element TEXT NOT NULL DEFAULT '', branch TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'available',
          description TEXT NOT NULL DEFAULT '', featured INTEGER NOT NULL DEFAULT 0, updated INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS orders (id TEXT PRIMARY KEY, sim_id INTEGER NOT NULL REFERENCES sims(id),
          number TEXT NOT NULL, price INTEGER NOT NULL, name TEXT NOT NULL, phone TEXT NOT NULL, address TEXT NOT NULL,
          branch TEXT NOT NULL, payment TEXT NOT NULL, paid INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'new',
          note TEXT NOT NULL DEFAULT '', staff_note TEXT NOT NULL DEFAULT '', assigned_to INTEGER REFERENCES users(id),
          lookup_hash TEXT NOT NULL, request_key TEXT UNIQUE NOT NULL, created INTEGER NOT NULL, updated INTEGER NOT NULL,
          hold_until INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS leads (id INTEGER PRIMARY KEY, name TEXT NOT NULL, phone TEXT NOT NULL, message TEXT NOT NULL,
          branch TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'new', staff_note TEXT NOT NULL DEFAULT '',
          assigned_to INTEGER REFERENCES users(id), follow_up TEXT NOT NULL DEFAULT '', created INTEGER NOT NULL, updated INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id),
          action TEXT NOT NULL, target TEXT NOT NULL, detail TEXT NOT NULL, created INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS limits (key TEXT PRIMARY KEY, count INTEGER NOT NULL, expires INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS metrics (day TEXT NOT NULL, page TEXT NOT NULL, count INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(day,page));
        CREATE TABLE IF NOT EXISTS outbox (id INTEGER PRIMARY KEY, subject TEXT NOT NULL, content TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', attempts INTEGER NOT NULL DEFAULT 0, next_try INTEGER NOT NULL, created INTEGER NOT NULL);
        CREATE INDEX IF NOT EXISTS orders_branch_status ON orders(branch,status);
        CREATE INDEX IF NOT EXISTS sims_status_branch ON sims(status,branch);
        ''')
        # Additive migration preserves existing consultation and historical order data.
        columns = {row['name'] for row in db.execute('PRAGMA table_info(leads)')}
        for name, declaration in {'sim_id': 'INTEGER REFERENCES sims(id)', 'sim_number': "TEXT NOT NULL DEFAULT ''", 'quoted_price': 'INTEGER', 'request_key': 'TEXT'}.items():
            if name not in columns:
                db.execute(f'ALTER TABLE leads ADD COLUMN {name} {declaration}')
        db.execute('CREATE UNIQUE INDEX IF NOT EXISTS leads_request_key ON leads(request_key)')
        db.execute('INSERT OR IGNORE INTO settings VALUES(1,?)', (json.dumps(DEFAULT_SETTINGS, ensure_ascii=False),))


def settings(db):
    return json.loads(db.execute('SELECT value FROM settings WHERE id=1').fetchone()[0])


def audit(db, action, target, detail):
    uid = g.user['id'] if getattr(g, 'user', None) else None
    db.execute('INSERT INTO audit(user_id,action,target,detail,created) VALUES(?,?,?,?,?)',
               (uid, action, str(target), json.dumps(detail, ensure_ascii=False), now()))


def expire_holds(db):
    rows = db.execute("SELECT id,sim_id FROM orders WHERE status='new' AND hold_until<=?", (now(),)).fetchall()
    for row in rows:
        db.execute("UPDATE orders SET status='expired',updated=? WHERE id=?", (now(), row['id']))
        db.execute("UPDATE sims SET status='available',updated=? WHERE id=? AND status='reserved'", (now(), row['sim_id']))
    db.execute('DELETE FROM sessions WHERE expires<?', (now(),))
    db.execute('DELETE FROM limits WHERE expires<?', (now(),))


def text(value, label, maximum=500, required=True):
    if not isinstance(value, str) or len(value.strip()) > maximum or (required and not value.strip()):
        raise Problem(f'{label} không hợp lệ (tối đa {maximum} ký tự).')
    return value.strip()


def phone(value):
    result = re.sub(r'[\s.+()-]', '', text(value, 'Số điện thoại', 30))
    if result.startswith('84'):
        result = '0' + result[2:]
    if not re.fullmatch(r'0[35789]\d{8}', result):
        raise Problem('Số điện thoại Việt Nam chưa hợp lệ.')
    return result


def integer(value, label, minimum=0, maximum=10**12):
    if isinstance(value, bool):
        raise Problem(f'{label} không hợp lệ.')
    try:
        result = int(value)
    except (ValueError, TypeError):
        raise Problem(f'{label} phải là số nguyên.')
    if str(value).strip() != str(result) or not minimum <= result <= maximum:
        raise Problem(f'{label} nằm ngoài giới hạn.')
    return result


def branch(value):
    if value not in ('cs1', 'cs2', 'cs3'):
        raise Problem('Vui lòng chọn cơ sở.')
    return value


def body():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise Problem('Yêu cầu phải là JSON.')
    return data


def limit(bucket, maximum, period):
    # Direct peer address is used unless a trusted reverse proxy is explicitly configured.
    key = hashlib.sha256(f'{bucket}:{request.remote_addr}'.encode()).hexdigest()
    with database(True) as db:
        row = db.execute('SELECT count,expires FROM limits WHERE key=?', (key,)).fetchone()
        if row and row['expires'] > now():
            if row['count'] >= maximum:
                raise Problem('Bạn thao tác quá nhanh. Vui lòng thử lại sau.', 429)
            db.execute('UPDATE limits SET count=count+1 WHERE key=?', (key,))
        else:
            db.execute('INSERT OR REPLACE INTO limits VALUES(?,1,?)', (key, now() + period))


def require_user(owner=False):
    if not getattr(g, 'user', None):
        raise Problem('Vui lòng đăng nhập.', 401)
    if owner and g.user['role'] != 'owner':
        raise Problem('Chức năng chỉ dành cho chủ cửa hàng.', 403)
    return g.user


def scope(db, table, key):
    user = require_user()
    row = db.execute(f'SELECT * FROM {table} WHERE id=?', (key,)).fetchone()
    if not row:
        raise Problem('Không tìm thấy dữ liệu.', 404)
    if user['role'] != 'owner' and row['branch'] != user['branch']:
        raise Problem('Bạn không được truy cập dữ liệu cơ sở khác.', 403)
    return row


def public_sim(row):
    return dict(row)


def valid_sim(data):
    number = phone(data.get('number'))
    kind = text(data.get('type'), 'Loại sim', 60)
    element = text(data.get('element', ''), 'Ngũ hành', 20, False)
    if element not in ('', 'Kim', 'Mộc', 'Thủy', 'Hỏa', 'Thổ'):
        raise Problem('Ngũ hành không hợp lệ.')
    status = data.get('status', 'available')
    if status not in ('available', 'hidden', 'reserved', 'sold'):
        raise Problem('Trạng thái kho không hợp lệ.')
    return {'number': number, 'price': integer(data.get('price'), 'Giá bán', 1), 'type': kind,
            'element': element, 'branch': branch(data.get('branch')), 'status': status,
            'description': text(data.get('description', ''), 'Mô tả', 2000, False),
            'featured': int(bool(data.get('featured', False))), 'updated': now()}


@app.before_request
def security():
    g.user = None
    g.session = None
    if request.path.startswith('/api/'):
        if request.method not in ('GET', 'HEAD', 'OPTIONS'):
            if not request.is_json:
                raise Problem('Chỉ chấp nhận JSON.', 415)
            origin = request.headers.get('Origin')
            site = os.environ.get('SITE_URL', '').rstrip('/')
            allowed = {request.host_url.rstrip('/')}
            if site:
                allowed.add(site)
            if origin and origin not in allowed:
                raise Problem('Nguồn yêu cầu không hợp lệ.', 403)
        token = request.cookies.get('session', '')
        if token:
            with database() as db:
                session = db.execute('SELECT * FROM sessions WHERE token=? AND expires>?',
                                     (hashlib.sha256(token.encode()).hexdigest(), now())).fetchone()
                if session:
                    row = db.execute('SELECT id,username,name,role,branch,active FROM users WHERE id=? AND active=1', (session['user_id'],)).fetchone()
                    if row:
                        g.user, g.session = dict(row), dict(session)
        if request.path.startswith('/api/admin/') and request.method not in ('GET', 'HEAD'):
            require_user()
            if not hmac.compare_digest(request.headers.get('X-CSRF-Token', ''), g.session['csrf']):
                raise Problem('Phiên thao tác không hợp lệ. Vui lòng tải lại trang.', 403)


@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Permissions-Policy'] = 'camera=(), microphone=(), geolocation=()'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    if request.path.startswith('/api/') or request.path == '/admin':
        response.headers['Cache-Control'] = 'no-store'
    if os.environ.get('SECURE_COOKIES') == '1':
        response.headers['Strict-Transport-Security'] = 'max-age=31536000'
    return response


@app.errorhandler(Problem)
def problem(error):
    return jsonify(error=error.message), error.status


@app.errorhandler(sqlite3.IntegrityError)
def integrity(error):
    return jsonify(error='Dữ liệu bị trùng hoặc không hợp lệ. Vui lòng kiểm tra lại.'), 409


@app.errorhandler(Exception)
def unexpected(error):
    if isinstance(error, HTTPException):
        return jsonify(error=error.description), error.code
    app.logger.exception('Request failed')
    return jsonify(error='Hệ thống gặp lỗi. Vui lòng thử lại hoặc gọi hotline.'), 500


def queue_mail(db, subject, content):
    db.execute('INSERT INTO outbox(subject,content,next_try,created) VALUES(?,?,?,?)', (subject, content, now(), now()))


@app.post('/api/view')
def record_view():
    limit('views', 120, 60)
    page = text(body().get('page'), 'Trang', 100)
    if page not in ('/', '/kho-sim', '/sim-phong-thuy', '/lien-he', '/chinh-sach', '/tra-cuu') and not re.fullmatch(r'/sim/\d{1,10}', page):
        raise Problem('Trang không hợp lệ.')
    from zoneinfo import ZoneInfo
    day = datetime.now(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%Y-%m-%d')
    with database(True) as db:
        db.execute('INSERT INTO metrics(day,page,count) VALUES(?,?,1) ON CONFLICT(day,page) DO UPDATE SET count=count+1', (day, page))
    return jsonify(ok=True)


@app.get('/api/health')
def health():
    with database() as db:
        db.execute('SELECT 1').fetchone()
    return jsonify(status='ok')


@app.get('/api/settings')
def get_settings():
    with database() as db:
        data = settings(db)
    return jsonify({k: v for k, v in data.items() if k not in {'bank_name', 'bank_account', 'bank_holder', 'hold_minutes', 'checkout_enabled', 'shipping_policy', 'return_policy'}})


@app.get('/api/sims')
def get_sims():
    clauses = ["status IN ('available','reserved','sold')"]
    args = []
    q = re.sub(r'\D', '', request.args.get('q', ''))
    if q:
        clauses.append('number LIKE ?')
        args.append('%' + q + '%')
    if request.args.get('status'):
        if request.args['status'] not in ('available', 'reserved', 'sold'):
            raise Problem('Trạng thái sim không hợp lệ.')
        clauses.append('status=?')
        args.append(request.args['status'])
    for field in ('type', 'element', 'branch'):
        if request.args.get(field):
            clauses.append(f'{field}=?')
            args.append(request.args[field])
    if request.args.get('prefix'):
        clauses.append('number LIKE ?')
        args.append(re.sub(r'\D', '', request.args['prefix']) + '%')
    for key, op in [('min_price', '>='), ('max_price', '<=')]:
        if request.args.get(key):
            clauses.append(f'price{op}?')
            args.append(integer(request.args[key], 'Khoảng giá'))
    page = integer(request.args.get('page', '1'), 'Trang', 1, 100000)
    size = integer(request.args.get('size', '24'), 'Số sản phẩm', 1, 60)
    sorting = {'asc': 'price ASC,id', 'desc': 'price DESC,id'}.get(request.args.get('sort'), 'featured DESC,id DESC')
    with database(True) as db:
        expire_holds(db)
        condition = ' AND '.join(clauses)
        total = db.execute(f'SELECT count(*) FROM sims WHERE {condition}', args).fetchone()[0]
        rows = db.execute(f'SELECT * FROM sims WHERE {condition} ORDER BY {sorting} LIMIT ? OFFSET ?', args + [size, (page - 1) * size]).fetchall()
    return jsonify(items=[public_sim(r) for r in rows], total=total, page=page, size=size)


@app.get('/api/sim-options')
def sim_options():
    with database() as db:
        types = [r[0] for r in db.execute("SELECT DISTINCT type FROM sims WHERE status!='hidden' ORDER BY type")]
    return jsonify(types=types)


@app.get('/api/sims/<int:sid>')
def get_sim(sid):
    with database(True) as db:
        expire_holds(db)
        row = db.execute("SELECT * FROM sims WHERE id=? AND status!='hidden'", (sid,)).fetchone()
        if not row:
            raise Problem('Không tìm thấy sim.', 404)
    return jsonify(public_sim(row))


@app.post('/api/orders')
@app.post('/api/orders/lookup')
def retired_checkout():
    raise Problem('Vui lòng chọn sim và gửi yêu cầu gọi tư vấn.', 410)


@app.post('/api/leads')
def create_lead():
    limit('leads', 8, 600)
    data = body()
    if data.get('website') or data.get('consent') is not True:
        raise Problem('Vui lòng đồng ý chính sách xử lý dữ liệu.')
    customer_name = text(data.get('name'), 'Họ tên', 100)
    customer_phone = phone(data.get('phone'))
    request_key = text(data.get('request_key', ''), 'Mã yêu cầu', 100, False)
    if request_key and not re.fullmatch(r'[a-zA-Z0-9_-]{20,100}', request_key):
        raise Problem('Mã yêu cầu không hợp lệ.')
    selected_id = integer(data['sim_id'], 'Mã sim', 1) if data.get('sim_id') is not None else None
    message = text(data.get('message', ''), 'Nội dung tư vấn', 2000, selected_id is None)
    with database(True) as db:
        if request_key:
            previous = db.execute('SELECT * FROM leads WHERE request_key=?', (request_key,)).fetchone()
            if previous:
                if (previous['name'], previous['phone'], previous['sim_id'], previous['message']) != (customer_name, customer_phone, selected_id, message):
                    raise Problem('Yêu cầu đã gửi có thông tin khác. Vui lòng tải lại trang.', 409)
                return jsonify(id=previous['id'], message='Đã nhận yêu cầu. Cửa hàng sẽ gọi tư vấn.'), 200
        if not settings(db).get('privacy_policy'):
            raise Problem('Cửa hàng chưa mở biểu mẫu tư vấn. Vui lòng gọi hotline.', 503)
        sim = None
        if selected_id is not None:
            sim = db.execute('SELECT * FROM sims WHERE id=?', (selected_id,)).fetchone()
            if not sim or sim['status'] != 'available':
                raise Problem('Sim này hiện không còn sẵn để tư vấn. Vui lòng chọn sim khác hoặc gọi cửa hàng.', 409)
        # Store an authoritative snapshot; consultation never reserves or sells stock.
        customer_branch = sim['branch'] if sim else branch(data.get('branch'))
        sid = db.execute('''INSERT INTO leads(name,phone,message,branch,created,updated,sim_id,sim_number,quoted_price,request_key)
                            VALUES(?,?,?,?,?,?,?,?,?,?)''',
                         (customer_name, customer_phone, message, customer_branch, now(), now(), selected_id,
                          sim['number'] if sim else '', sim['price'] if sim else None, request_key or None)).lastrowid
        queue_mail(db, 'Khách cần gọi tư vấn', 'Có yêu cầu mới. Vào trang quản trị để xem sim khách chọn và gọi lại.')
    return jsonify(id=sid, message='Đã nhận yêu cầu. Cửa hàng sẽ gọi tư vấn.'), 201


@app.post('/api/login')
def login():
    limit('login', 10, 900)
    data = body()
    username = text(data.get('username'), 'Tên đăng nhập', 60)
    password = text(data.get('password'), 'Mật khẩu', 200)
    with database(True) as db:
        row = db.execute('SELECT * FROM users WHERE username=? AND active=1', (username,)).fetchone()
        # Perform password hashing even for unknown names to reduce timing differences.
        password_hash = row['password'] if row else generate_password_hash('invalid-placeholder')
        if not check_password_hash(password_hash, password) or not row:
            raise Problem('Tên đăng nhập hoặc mật khẩu không đúng.', 401)
        token, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
        db.execute('INSERT INTO sessions VALUES(?,?,?,?)', (hashlib.sha256(token.encode()).hexdigest(), row['id'], csrf, now() + 28800))
        g.user = dict(row)
        audit(db, 'auth.login', row['id'], {})
    response = jsonify(user={k: row[k] for k in ('id', 'username', 'name', 'role', 'branch')}, csrf=csrf)
    response.set_cookie('session', token, httponly=True, secure=os.environ.get('SECURE_COOKIES') == '1', samesite='Strict', max_age=28800, path='/')
    return response


@app.get('/api/admin/me')
def me():
    return jsonify(user=require_user(), csrf=g.session['csrf'])


@app.post('/api/admin/logout')
def logout():
    with database(True) as db:
        db.execute('DELETE FROM sessions WHERE token=?', (g.session['token'],))
    response = jsonify(ok=True)
    response.delete_cookie('session', path='/')
    return response


@app.post('/api/admin/password')
def password_change():
    user = require_user()
    data = body()
    new = text(data.get('new_password'), 'Mật khẩu mới', 200)
    if len(new) < 12:
        raise Problem('Mật khẩu cần ít nhất 12 ký tự.')
    with database(True) as db:
        row = db.execute('SELECT password FROM users WHERE id=?', (user['id'],)).fetchone()
        if not check_password_hash(row['password'], data.get('current_password', '')):
            raise Problem('Mật khẩu hiện tại không đúng.', 403)
        db.execute('UPDATE users SET password=? WHERE id=?', (generate_password_hash(new), user['id']))
        db.execute('DELETE FROM sessions WHERE user_id=?', (user['id'],))
        audit(db, 'user.password_changed', user['id'], {})
    response = jsonify(ok=True)
    response.delete_cookie('session', path='/')
    return response


@app.get('/api/admin/dashboard')
def dashboard():
    user = require_user()
    with database(True) as db:
        expire_holds(db)
        args = [] if user['role'] == 'owner' else [user['branch']]
        where = '' if not args else ' WHERE branch=?'
        counts = {}
        for table in ('sims', 'leads'):
            counts[table] = [dict(r) for r in db.execute(f'SELECT status,count(*) AS count FROM {table}{where} GROUP BY status', args)]
        counts['views'] = [dict(r) for r in db.execute('SELECT day,sum(count) AS count FROM metrics GROUP BY day ORDER BY day DESC LIMIT 7')] if user['role'] == 'owner' else []
        counts['popular_pages'] = [dict(r) for r in db.execute('SELECT page,sum(count) AS count FROM metrics GROUP BY page ORDER BY count DESC LIMIT 10')] if user['role'] == 'owner' else []
        counts['mail'] = [dict(r) for r in db.execute('SELECT status,count(*) AS count FROM outbox GROUP BY status')] if user['role'] == 'owner' else []
    return jsonify(counts)


@app.get('/api/admin/<table>')
def admin_list(table):
    user = require_user()
    if table not in ('sims', 'leads', 'audit', 'users'):
        raise Problem('Không tìm thấy.', 404)
    if table in ('audit', 'users'):
        require_user(True)
    page = integer(request.args.get('page', '1'), 'Trang', 1, 100000)
    size = integer(request.args.get('size', '30'), 'Số dòng', 1, 100)
    conditions, args = [], []
    if table in ('sims', 'orders', 'leads'):
        if user['role'] != 'owner':
            conditions.append('branch=?')
            args.append(user['branch'])
        elif request.args.get('branch'):
            conditions.append('branch=?')
            args.append(branch(request.args['branch']))
        if request.args.get('status'):
            conditions.append('status=?')
            args.append(request.args['status'])
        if request.args.get('q'):
            q = '%' + request.args['q'][:100] + '%'
            fields = ['number'] if table == 'sims' else ['name', 'phone', 'id'] + (['sim_number'] if table == 'leads' else [])
            conditions.append('(' + ' OR '.join(f'{f} LIKE ?' for f in fields) + ')')
            args.extend([q] * len(fields))
    where = ' WHERE ' + ' AND '.join(conditions) if conditions else ''
    fields = 'id,username,name,role,branch,active' if table == 'users' else '*'
    with database(True) as db:
        expire_holds(db)
        total = db.execute(f'SELECT count(*) FROM {table}{where}', args).fetchone()[0]
        rows = db.execute(f'SELECT {fields} FROM {table}{where} ORDER BY id DESC LIMIT ? OFFSET ?', args + [size, (page - 1) * size]).fetchall()
    items = [dict(r) for r in rows]
    for item in items:
        item.pop('lookup_hash', None)
        item.pop('request_key', None)
    return jsonify(items=items, total=total, page=page, size=size)


@app.post('/api/admin/sims')
def add_sim():
    require_user(True)
    data = valid_sim(body())
    with database(True) as db:
        sid = db.execute('INSERT INTO sims(number,price,type,element,branch,status,description,featured,updated) VALUES(:number,:price,:type,:element,:branch,:status,:description,:featured,:updated)', data).lastrowid
        audit(db, 'sim.created', sid, data)
    return jsonify(id=sid), 201


@app.patch('/api/admin/sims/<int:sid>')
def edit_sim(sid):
    require_user(True)
    data = valid_sim(body())
    with database(True) as db:
        old = scope(db, 'sims', sid)
        if old['number'] != data['number'] and (db.execute('SELECT 1 FROM orders WHERE sim_id=? LIMIT 1', (sid,)).fetchone() or db.execute('SELECT 1 FROM leads WHERE sim_id=? LIMIT 1', (sid,)).fetchone()):
            raise Problem('Không đổi số của sim đã có lịch sử tư vấn hoặc giao dịch.', 409)
        db.execute('UPDATE sims SET number=:number,price=:price,type=:type,element=:element,branch=:branch,status=:status,description=:description,featured=:featured,updated=:updated WHERE id=:id', {**data, 'id': sid})
        audit(db, 'sim.updated', sid, {'before': dict(old), 'after': data})
    return jsonify(ok=True)


@app.post('/api/admin/sims/bulk-price')
def bulk_price():
    require_user(True)
    data = body()
    ids = data.get('ids')
    if not isinstance(ids, list) or not 1 <= len(ids) <= 500:
        raise Problem('Chọn từ 1 đến 500 sim.')
    ids = list(set(integer(x, 'Mã sim', 1) for x in ids))
    new_price = integer(data.get('price'), 'Giá bán', 1)
    with database(True) as db:
        rows = db.execute('SELECT * FROM sims WHERE id IN (' + ','.join('?' for _ in ids) + ')', ids).fetchall()
        if len(rows) != len(ids) or any(r['status'] in ('reserved', 'sold') for r in rows):
            raise Problem('Có sim không tồn tại, đang giữ chỗ hoặc đã bán.', 409)
        db.execute('UPDATE sims SET price=?,updated=? WHERE id IN (' + ','.join('?' for _ in ids) + ')', [new_price, now()] + ids)
        audit(db, 'sim.bulk_price', 'batch', {'before': [{'id': r['id'], 'price': r['price']} for r in rows], 'price': new_price})
    return jsonify(count=len(ids))


@app.post('/api/admin/sims/import')
def import_sims():
    require_user(True)
    return import_content(text(body().get('csv'), 'CSV', 1024 * 1024))


def import_content(content):
    reader = csv.DictReader(io.StringIO(content.lstrip('\ufeff')))
    if not reader.fieldnames or not {'number', 'price', 'type', 'branch'}.issubset(reader.fieldnames):
        raise Problem('CSV cần các cột number,price,type,branch. Số sim cần giữ số 0 đầu.')
    rows, seen = [], set()
    for line, row in enumerate(reader, 2):
        if len(rows) >= 5000:
            raise Problem('Mỗi lần nhập tối đa 5.000 sim.')
        try:
            row['featured'] = row.get('featured', '') in ('1', 'true')
            row['status'] = row.get('status') or 'available'
            parsed = valid_sim(row)
        except Problem as error:
            raise Problem(f'Dòng {line}: {error.message}')
        if parsed['number'] in seen:
            raise Problem(f'Dòng {line}: số sim trùng trong tệp.')
        seen.add(parsed['number'])
        rows.append(parsed)
    if not rows:
        raise Problem('Tệp không có sim.')
    with database(True) as db:
        for row in rows:
            if db.execute('SELECT 1 FROM sims WHERE number=?', (row['number'],)).fetchone():
                raise Problem(f"Số {row['number']} đã có trong kho. Chưa nhập dòng nào.", 409)
        db.executemany('INSERT INTO sims(number,price,type,element,branch,status,description,featured,updated) VALUES(:number,:price,:type,:element,:branch,:status,:description,:featured,:updated)', rows)
        audit(db, 'sim.import', 'batch', {'count': len(rows)})
    return jsonify(count=len(rows)), 201


def csv_safe(value):
    value = str(value if value is not None else '')
    return "'" + value if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


@app.get('/api/admin/export/<table>')
def export(table):
    user = require_user()
    if table not in ('sims', 'orders', 'leads'):
        raise Problem('Không tìm thấy.', 404)
    where, args = ('', []) if user['role'] == 'owner' else (' WHERE branch=?', [user['branch']])
    with database() as db:
        rows = db.execute(f'SELECT * FROM {table}{where} ORDER BY id DESC', args).fetchall()
        columns = [r[1] for r in db.execute(f'PRAGMA table_info({table})') if r[1] not in ('lookup_hash', 'request_key')]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([csv_safe(row[k]) for k in columns])
    return Response('\ufeff' + output.getvalue(), content_type='text/csv; charset=utf-8', headers={'Content-Disposition': f'attachment; filename="{table}.csv"'})


def assignment(db, value, owned_branch, current=None):
    if value in (None, ''):
        if g.user['role'] != 'owner' and current not in (None, g.user['id']):
            raise Problem('Không được bỏ phân công của nhân viên khác.', 403)
        return None
    uid = integer(value, 'Nhân viên', 1)
    row = db.execute('SELECT * FROM users WHERE id=? AND active=1', (uid,)).fetchone()
    if not row or (row['role'] != 'owner' and row['branch'] != owned_branch):
        raise Problem('Nhân viên không thuộc cơ sở phụ trách.')
    if g.user['role'] != 'owner' and uid not in (g.user['id'], current):
        raise Problem('Nhân viên chỉ được nhận việc cho chính mình.', 403)
    return uid


@app.patch('/api/admin/leads/<int:lid>')
def update_lead(lid):
    require_user()
    data = body()
    with database(True) as db:
        old = scope(db, 'leads', lid)
        status = data.get('status', old['status'])
        if status not in ('new', 'contacted', 'done'):
            raise Problem('Trạng thái tư vấn không hợp lệ.')
        follow_up = text(data.get('follow_up', old['follow_up']), 'Lịch gọi lại', 30, False)
        if follow_up:
            try:
                datetime.fromisoformat(follow_up)
            except ValueError:
                raise Problem('Lịch gọi lại không hợp lệ.')
        note = text(data.get('staff_note', old['staff_note']), 'Ghi chú nội bộ', 2000, False)
        assigned = assignment(db, data.get('assigned_to', old['assigned_to']), old['branch'], old['assigned_to'])
        db.execute('UPDATE leads SET status=?,staff_note=?,assigned_to=?,follow_up=?,updated=? WHERE id=?', (status, note, assigned, follow_up, now(), lid))
        audit(db, 'lead.updated', lid, {'status': status, 'assigned_to': assigned, 'follow_up': follow_up})
    return jsonify(ok=True)


@app.patch('/api/admin/settings')
def update_settings():
    require_user(True)
    data = body()
    allowed = set(DEFAULT_SETTINGS)
    if set(data) - allowed:
        raise Problem('Trường cấu hình không được hỗ trợ.')
    with database(True) as db:
        config = settings(db)
        for key, value in data.items():
            if key == 'branches':
                continue
            maximum = 10000 if key.endswith('_policy') else 500
            config[key] = text(value, key, maximum, key in ('name', 'phone', 'email'))
        phone(config['phone'])
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', config['email']):
            raise Problem('Email cửa hàng không hợp lệ.')
        for key in ('site_url', 'zalo_url'):
            if config[key]:
                url = urlsplit(config[key])
                if url.scheme != 'https' or not url.hostname or url.username or url.password:
                    raise Problem(f'{key} cần là đường dẫn HTTPS hợp lệ.')
                if key == 'zalo_url' and url.hostname not in ('zalo.me', 'oa.zalo.me'):
                    raise Problem('Link Zalo cần thuộc zalo.me hoặc oa.zalo.me.')
                if key == 'site_url' and (url.path not in ('', '/') or url.query or url.fragment):
                    raise Problem('Tên miền không chứa đường dẫn, tham số hoặc fragment.')
        if config['banner_url'] and not re.fullmatch(r'/uploads/[a-f0-9]{32}\.(png|jpg|webp)', config['banner_url']):
            raise Problem('Banner phải là ảnh được tải lên bằng trang quản trị.')
        if 'branches' in data:
            if not isinstance(data['branches'], list) or len(data['branches']) != 3:
                raise Problem('Cần đủ 3 cơ sở.')
            branches = []
            for index, entry in enumerate(data['branches'], 1):
                if not isinstance(entry, dict) or entry.get('id') != f'cs{index}':
                    raise Problem('Mã cơ sở không hợp lệ.')
                map_url = text(entry.get('map_url', ''), 'Link bản đồ', 500, False)
                if map_url:
                    url = urlsplit(map_url)
                    if url.scheme != 'https' or url.hostname not in ('www.google.com', 'maps.google.com', 'maps.app.goo.gl', 'goo.gl') or url.username or url.password:
                        raise Problem('Link bản đồ phải là link Google Maps HTTPS.')
                branches.append({'id': f'cs{index}', 'name': f'CS{index}', 'address': text(entry.get('address'), 'Địa chỉ', 300), 'map_url': map_url})
            config['branches'] = branches
        db.execute('UPDATE settings SET value=? WHERE id=1', (json.dumps(config, ensure_ascii=False),))
        audit(db, 'settings.updated', 'store', {'fields': list(data)})
    return jsonify(ok=True)


@app.post('/api/admin/uploads')
def upload():
    require_user(True)
    import base64
    data = body()
    try:
        raw = base64.b64decode(data.get('content', ''), validate=True)
    except (ValueError, TypeError):
        raise Problem('Tệp ảnh không hợp lệ.')
    if not 16 <= len(raw) <= 1024 * 1024:
        raise Problem('Ảnh cần nhỏ hơn 1 MB.')
    from PIL import Image, UnidentifiedImageError
    Image.MAX_IMAGE_PIXELS = 16_000_000
    try:
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('PNG', 'JPEG', 'WEBP') or image.width * image.height > 16_000_000:
                raise Problem('Ảnh không hợp lệ hoặc quá lớn.')
            image.load()
            output = io.BytesIO()
            image.convert('RGB').save(output, format='JPEG', quality=88)
            raw = output.getvalue()
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError, ValueError):
        raise Problem('Không đọc được ảnh PNG, JPEG hoặc WebP hợp lệ.')
    ext = 'jpg'
    folder = DATA / 'uploads'
    folder.mkdir(exist_ok=True)
    filename = secrets.token_hex(16) + '.' + ext
    (folder / filename).write_bytes(raw)
    with database(True) as db:
        audit(db, 'image.uploaded', filename, {'bytes': len(raw)})
    return jsonify(url='/uploads/' + filename), 201


@app.post('/api/admin/users')
def add_user():
    require_user(True)
    data = body()
    username = text(data.get('username'), 'Tên đăng nhập', 60)
    if not re.fullmatch(r'[a-zA-Z0-9._-]{3,60}', username):
        raise Problem('Tên đăng nhập gồm 3–60 chữ/số, dấu chấm, gạch ngang hoặc gạch dưới.')
    password = text(data.get('password'), 'Mật khẩu', 200)
    if len(password) < 12:
        raise Problem('Mật khẩu cần ít nhất 12 ký tự.')
    role = data.get('role', 'staff')
    if role not in ('owner', 'staff'):
        raise Problem('Vai trò không hợp lệ.')
    b = None if role == 'owner' else branch(data.get('branch'))
    with database(True) as db:
        uid = db.execute('INSERT INTO users(username,name,password,role,branch) VALUES(?,?,?,?,?)', (username, text(data.get('name'), 'Tên nhân viên', 100), generate_password_hash(password), role, b)).lastrowid
        audit(db, 'user.created', uid, {'username': username, 'role': role, 'branch': b})
    return jsonify(id=uid), 201


@app.patch('/api/admin/users/<int:uid>')
def edit_user(uid):
    user = require_user(True)
    data = body()
    if uid == user['id']:
        raise Problem('Không khóa hoặc thay quyền tài khoản đang dùng.')
    with database(True) as db:
        row = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
        if not row:
            raise Problem('Không tìm thấy nhân viên.', 404)
        active = data.get('active', bool(row['active']))
        if not isinstance(active, bool):
            raise Problem('Trạng thái tài khoản không hợp lệ.')
        role = data.get('role', row['role'])
        if role not in ('owner', 'staff'):
            raise Problem('Vai trò không hợp lệ.')
        b = None if role == 'owner' else branch(data.get('branch', row['branch']))
        password = row['password']
        if data.get('password'):
            new = text(data['password'], 'Mật khẩu', 200)
            if len(new) < 12:
                raise Problem('Mật khẩu cần ít nhất 12 ký tự.')
            password = generate_password_hash(new)
        db.execute('UPDATE users SET name=?,role=?,branch=?,active=?,password=? WHERE id=?', (text(data.get('name', row['name']), 'Tên nhân viên', 100), role, b, int(active), password, uid))
        db.execute('DELETE FROM sessions WHERE user_id=?', (uid,))
        audit(db, 'user.updated', uid, {'active': active, 'role': role, 'branch': b})
    return jsonify(ok=True)


@app.get('/api/admin/backup')
def backup_download():
    require_user(True)
    # A consistent SQLite snapshot includes users and customer data; never serve the live DB.
    temp = DATA / ('backup-' + secrets.token_hex(12) + '.sqlite3')
    with database() as source:
        with sqlite3.connect(temp) as destination:
            source.backup(destination)
            destination.execute('PRAGMA journal_mode=DELETE')
            destination.execute('DELETE FROM sessions')
            destination.execute('DELETE FROM limits')
    destination.close()
    os.chmod(temp, 0o600)
    content = temp.read_bytes()
    temp.unlink()
    return Response(content, content_type='application/octet-stream', headers={'Content-Disposition': 'attachment; filename="viettel-backup.sqlite3"'})


@app.get('/uploads/<filename>')
def uploaded_file(filename):
    if not re.fullmatch(r'[a-f0-9]{32}\.(png|jpg|webp)', filename):
        raise Problem('Không tìm thấy ảnh.', 404)
    return send_file(DATA / 'uploads' / filename)


def public_meta(sid=None):
    with database() as db:
        config = settings(db)
        sim = db.execute("SELECT * FROM sims WHERE id=? AND status!='hidden'", (sid,)).fetchone() if sid else None
    title = f"Sim {sim['number']} | {config['name']}" if sim else f"{config['name']} | Sim số đẹp tại Phú Thọ"
    description = sim['description'] or f"Sim Viettel {sim['number']} — {sim['price']:,} đồng. Liên hệ {config['phone']}." if sim else f"Kho sim số đẹp Viettel. Liên hệ {config['phone']}."
    url = config['site_url'].rstrip('/')
    return config, sim, title, description, url


@app.get('/tra-cuu')
def retired_tracking():
    return redirect('/lien-he', code=302)


@app.get('/')
@app.get('/kho-sim')
@app.get('/sim-phong-thuy')
@app.get('/lien-he')
@app.get('/chinh-sach')
@app.get('/sim/<int:sid>')
def index(sid=None):
    import html
    config, sim, title, description, url = public_meta(sid)
    if sid and not sim:
        return Response('Không tìm thấy sim.', status=404, content_type='text/plain; charset=utf-8')
    page = (ROOT / 'index.html').read_text()
    page = page.replace('<title>Viettel Phú Hà | Sim số đẹp</title>', '<title>' + html.escape(title) + '</title>')
    page = re.sub(r'<meta name="description"[^>]*>', '<meta name="description" content="' + html.escape(description, quote=True) + '">', page)
    meta = f'<meta property="og:title" content="{html.escape(title, quote=True)}"><meta property="og:description" content="{html.escape(description, quote=True)}"><meta property="og:type" content="website">'
    if url:
        canonical = url + request.path
        meta += f'<link rel="canonical" href="{html.escape(canonical, quote=True)}"><meta property="og:url" content="{html.escape(canonical, quote=True)}"><meta property="og:image" content="{html.escape(url + (config["banner_url"] or "/assets/social.svg"), quote=True)}">'
    schema = {'@context': 'https://schema.org', '@type': 'Store', 'name': config['name'], 'telephone': config['phone'], 'email': config['email'], 'url': url}
    if sim:
        schema = {'@context': 'https://schema.org', '@type': 'Product', 'name': f"Sim {sim['number']}", 'description': description, 'offers': {'@type': 'Offer', 'price': sim['price'], 'priceCurrency': 'VND', 'availability': 'https://schema.org/InStock' if sim['status'] == 'available' else 'https://schema.org/OutOfStock'}}
    meta += '<script type="application/ld+json">' + json.dumps(schema, ensure_ascii=False).replace('<', '\\u003c') + '</script>'
    return Response(page.replace('</head>', meta + '</head>'), content_type='text/html; charset=utf-8')


@app.get('/sitemap.xml')
def sitemap():
    import html
    with database() as db:
        config = settings(db)
        rows = db.execute("SELECT id FROM sims WHERE status!='hidden'").fetchall()
    base = config['site_url'].rstrip('/')
    if not base:
        return Response('Tên miền chưa được cấu hình.', status=503)
    paths = ['/', '/kho-sim', '/sim-phong-thuy', '/lien-he', '/chinh-sach'] + ['/sim/' + str(r['id']) for r in rows]
    xml = '<?xml version="1.0" encoding="UTF-8"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">' + ''.join('<url><loc>' + html.escape(base + p) + '</loc></url>' for p in paths) + '</urlset>'
    return Response(xml, content_type='application/xml; charset=utf-8')


@app.get('/robots.txt')
def robots():
    with database() as db:
        base = settings(db)['site_url'].rstrip('/')
    return Response('User-agent: *\nDisallow: /admin\nDisallow: /api/\nDisallow: /tra-cuu\n' + ('Sitemap: ' + base + '/sitemap.xml\n' if base else ''), content_type='text/plain')


@app.get('/admin')
def admin_page():
    return send_file(ROOT / 'admin.html')


@app.get('/<filename>')
def static_file(filename):
    if filename not in ('app.js', 'admin.js', 'styles.css', 'admin.css', 'favicon.svg'):
        raise Problem('Không tìm thấy.', 404)
    return send_file(ROOT / filename)


@app.get('/assets/<filename>')
def assets(filename):
    if filename not in ('sim.svg', 'landscape.svg', 'social.svg'):
        raise Problem('Không tìm thấy.', 404)
    return send_file(ROOT / 'assets' / filename)


def xlsx_bytes(columns, rows):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Du lieu')
    sheet.append(columns)
    for row in rows:
        cells = []
        for column in columns:
            value = row[column] if column in row.keys() else ''
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            cells.append(cell)
        sheet.append(cells)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@app.get('/api/admin/template.xlsx')
def excel_template():
    require_user(True)
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Kho sim'
    sheet.append(['number', 'price', 'type', 'branch', 'element', 'description', 'featured', 'status'])
    for row in sheet.iter_rows(min_row=2, max_row=5001, min_col=1, max_col=1):
        row[0].number_format = '@'
    guide = book.create_sheet('Huong dan')
    for line in ['Điền dữ liệu tại sheet Kho sim; chỉ sheet đầu tiên được nhập.',
                 'number: số điện thoại 10 số, giữ số 0 đầu; cột này đã định dạng văn bản.',
                 'price: giá bán nguyên đồng, lớn hơn 0.',
                 'type: ví dụ Sim số đẹp, Sim tứ quý, Sim lộc phát.',
                 'branch: cs1 / cs2 / cs3.', 'element: để trống hoặc Kim, Mộc, Thủy, Hỏa, Thổ.',
                 'featured: 1 để làm nổi bật, 0 hoặc trống nếu không.', 'status: available (còn hàng) hoặc hidden (ẩn).',
                 'Không dùng công thức. Tối đa 5.000 dòng. Dòng sai hoặc số trùng sẽ hủy toàn bộ lần nhập.']:
        guide.append([line])
    output = io.BytesIO()
    book.save(output)
    return Response(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="mau-kho-sim.xlsx"'})


@app.post('/api/admin/sims/import-xlsx')
def import_excel():
    require_user(True)
    import base64
    import zipfile
    from openpyxl import load_workbook
    try:
        content = base64.b64decode(body().get('content', ''), validate=True)
        if len(content) > 1024 * 1024:
            raise Problem('Excel cần nhỏ hơn 1 MB.')
        with zipfile.ZipFile(io.BytesIO(content)) as zipped:
            if len(zipped.infolist()) > 1000 or sum(i.file_size for i in zipped.infolist()) > 20 * 1024 * 1024:
                raise Problem('Tệp Excel giải nén quá lớn.')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column > 20 or sheet.max_row > 10000:
                raise Problem('Excel cần tối đa 20 cột và 5.000 sim.')
            output = io.StringIO()
            writer = csv.writer(output)
            for row in sheet.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise Problem('Excel có công thức. Hãy dán giá trị trước khi nhập.')
                values = [cell.value if cell.value is not None else '' for cell in row]
                if not any(str(v).strip() for v in values):
                    continue
                writer.writerow(values)
        finally:
            book.close()
    except Problem:
        raise
    except Exception:
        raise Problem('Không đọc được Excel. Dùng tệp .xlsx theo mẫu.')
    return import_content(output.getvalue())


@app.get('/api/admin/export/<table>.xlsx')
def export_excel(table):
    user = require_user()
    if table not in ('sims', 'orders', 'leads'):
        raise Problem('Không tìm thấy.', 404)
    where, args = ('', []) if user['role'] == 'owner' else (' WHERE branch=?', [user['branch']])
    with database() as db:
        rows = db.execute(f'SELECT * FROM {table}{where} ORDER BY id DESC', args).fetchall()
        columns = [r[1] for r in db.execute(f'PRAGMA table_info({table})') if r[1] not in ('lookup_hash', 'request_key')]
    return Response(xlsx_bytes(columns, rows), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="{table}.xlsx"'})


initialize()


def deliver_mail():
    """Run one worker only. SMTP is optional; the admin dashboard works without it."""
    import smtplib
    import ssl
    from email.message import EmailMessage
    host = os.environ.get('SMTP_HOST', '')
    sender = os.environ.get('SMTP_FROM', '')
    if not host or not sender:
        return
    try:
        port = int(os.environ.get('SMTP_PORT', '465'))
        if port not in (465, 587):
            raise ValueError('SMTP requires TLS on port 465 or 587')
        with database() as db:
            rows = db.execute("SELECT * FROM outbox WHERE status='pending' AND next_try<=? ORDER BY id LIMIT 20", (now(),)).fetchall()
            config = settings(db)
        if not rows:
            return
        context = ssl.create_default_context()
        smtp = smtplib.SMTP_SSL(host, port, timeout=15, context=context) if port == 465 else smtplib.SMTP(host, port, timeout=15)
        with smtp:
            if port == 587:
                smtp.ehlo()
                smtp.starttls(context=context)
                smtp.ehlo()
            if os.environ.get('SMTP_USERNAME'):
                smtp.login(os.environ['SMTP_USERNAME'], os.environ.get('SMTP_PASSWORD', ''))
            for row in rows:
                message = EmailMessage()
                message['From'] = sender
                message['To'] = config['email']
                message['Subject'] = row['subject']
                message.set_content(row['content'] + ('\n' + config['site_url'].rstrip('/') + '/admin' if config['site_url'] else ''))
                try:
                    smtp.send_message(message)
                    with database(True) as db:
                        db.execute("UPDATE outbox SET status='sent',attempts=attempts+1 WHERE id=?", (row['id'],))
                except Exception:
                    retry_mail(row)
    except Exception:
        # Do not print SMTP credentials, server replies or customer information.
        app.logger.warning('Email delivery unavailable; check SMTP configuration. No credentials logged.')
        if 'rows' in locals():
            for row in rows:
                retry_mail(row)


def retry_mail(row):
    with database(True) as db:
        db.execute("UPDATE outbox SET attempts=attempts+1,next_try=? WHERE id=? AND status='pending'", (now() + min(3600, 60 * 2 ** min(row['attempts'], 6)), row['id']))


def cli():
    parser = argparse.ArgumentParser()
    parser.add_argument('command', choices=['serve', 'create-admin', 'backup', 'restore', 'mail-worker'])
    parser.add_argument('--file')
    parser.add_argument('--once', action='store_true')
    parser.add_argument('--port', type=int, default=int(os.environ.get('PORT', '3000')))
    args = parser.parse_args()
    if args.command == 'mail-worker':
        if args.once:
            deliver_mail()
        else:
            while True:
                deliver_mail()
                time.sleep(30)
    elif args.command == 'serve':
        app.run(host='0.0.0.0', port=args.port, debug=False)
    elif args.command == 'create-admin':
        import getpass
        username = text(input('Tên đăng nhập: '), 'Tên đăng nhập', 60)
        if not re.fullmatch(r'[a-zA-Z0-9._-]{3,60}', username):
            raise SystemExit('Tên đăng nhập không hợp lệ.')
        password = getpass.getpass('Mật khẩu (ít nhất 12 ký tự): ')
        if len(password) < 12 or password != getpass.getpass('Nhập lại mật khẩu: '):
            raise SystemExit('Mật khẩu không hợp lệ hoặc không khớp.')
        with database(True) as db:
            db.execute('INSERT INTO users(username,name,password,role) VALUES(?,?,?,?)', (username, 'Chủ cửa hàng', generate_password_hash(password), 'owner'))
        print('Đã tạo chủ cửa hàng. Đăng nhập tại /admin; không có mật khẩu mặc định.')
    elif args.command == 'backup':
        if not args.file:
            raise SystemExit('Cần --file trỏ tới tệp sao lưu ngoài thư mục công khai.')
        destination = Path(args.file)
        if destination.exists():
            raise SystemExit('Tệp đã tồn tại; không ghi đè.')
        with database() as source, sqlite3.connect(destination) as target:
            source.backup(target)
            target.execute('PRAGMA journal_mode=DELETE')
            target.execute('DELETE FROM sessions')
            target.execute('DELETE FROM limits')
        target.close()
        os.chmod(destination, 0o600)
        print('Đã sao lưu SQLite. Cần sao lưu riêng thư mục uploads.')
    else:
        if not args.file or not Path(args.file).is_file():
            raise SystemExit('Cần --file tới bản sao lưu hợp lệ. Dừng máy chủ trước khi khôi phục.')
        source_path = Path(args.file).resolve()
        if source_path == DB:
            raise SystemExit('Không khôi phục từ chính cơ sở dữ liệu đang dùng.')
        source = sqlite3.connect('file:' + str(source_path) + '?mode=ro', uri=True)
        try:
            if source.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
                raise SystemExit('Bản sao lưu hỏng.')
            tables = {r[0] for r in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if not {'settings', 'users', 'sessions', 'sims', 'orders', 'leads', 'audit', 'limits', 'metrics', 'outbox'} <= tables:
                raise SystemExit('Bản sao lưu không thuộc hệ thống này.')
            safety = DATA / ('before-restore-' + str(now()) + '.sqlite3')
            with database() as current, sqlite3.connect(safety) as target:
                current.backup(target)
                target.execute('PRAGMA journal_mode=DELETE')
            with sqlite3.connect(DB) as destination:
                source.backup(destination)
                destination.execute('DELETE FROM sessions')
                destination.execute('DELETE FROM limits')
            destination.close()
            target.close()
            os.chmod(safety, 0o600)
            print('Đã khôi phục; bản trước khôi phục được giữ trong DATA_DIR. Khởi động lại máy chủ.')
        finally:
            source.close()


if __name__ == '__main__':
    cli()
