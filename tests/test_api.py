import base64
import csv
import io
import json
import os
import secrets
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

# All fixtures are isolated from the real store; no test credentials are published.
temporary = tempfile.TemporaryDirectory(prefix='viettel-api-')
os.environ['DATA_DIR'] = temporary.name
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server
from werkzeug.security import generate_password_hash
from openpyxl import Workbook, load_workbook
from PIL import Image


class StoreTests(unittest.TestCase):
    def setUp(self):
        for path in server.DATA.glob('store.db*'):
            path.unlink()
        server.initialize()
        self.password = secrets.token_urlsafe(24)
        with server.database(True) as db:
            for username, role, branch in [('owner', 'owner', None), ('staff1', 'staff', 'cs1'), ('staff2', 'staff', 'cs2')]:
                db.execute('INSERT INTO users(username,name,password,role,branch) VALUES(?,?,?,?,?)',
                           (username, username, generate_password_hash(self.password), role, branch))
        self.client = server.app.test_client()
        self.csrf = self.login(self.client, 'owner')
        self.sid = self.add_sim('0344639666')['id']

    def login(self, client, username):
        response = client.post('/api/login', json={'username': username, 'password': self.password})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertIn('HttpOnly', response.headers['Set-Cookie'])
        self.assertIn('SameSite=Strict', response.headers['Set-Cookie'])
        return response.json['csrf']

    def mutate(self, path, data, method='post', client=None, csrf=None):
        return getattr(client or self.client, method)(path, json=data, headers={'X-CSRF-Token': csrf or self.csrf})

    def add_sim(self, number, branch='cs1'):
        response = self.mutate('/api/admin/sims', {'number': number, 'price': 1500000, 'type': 'Sim số đẹp', 'branch': branch})
        self.assertEqual(response.status_code, 201, response.json)
        return response.json

    def open_checkout(self):
        response = self.mutate('/api/admin/settings', {'hours': '08:00–20:00', 'privacy_policy': 'Chỉ sử dụng thông tin để xử lý yêu cầu.',
                           'shipping_policy': 'Xác nhận phí trước khi giao.', 'return_policy': 'Liên hệ cửa hàng để xử lý đổi trả.',
                           'registration_policy': 'Đăng ký theo quy định nhà mạng.', 'checkout_enabled': True}, 'patch')
        self.assertEqual(response.status_code, 200, response.json)

    def order_data(self, sid=None):
        return {'sim_id': sid or self.sid, 'name': 'Khách kiểm thử', 'phone': '0988123456', 'address': 'Địa chỉ thử nghiệm',
                'branch': 'cs2', 'payment': 'cod', 'consent': True, 'request_key': secrets.token_hex(20), 'lookup_token': secrets.token_hex(20)}

    def create_order(self):
        self.open_checkout()
        data = self.order_data()
        data['price'] = 1
        response = server.app.test_client().post('/api/orders', json=data)
        self.assertEqual(response.status_code, 201, response.json)
        return response.json, data

    def test_empty_store_and_checkout_gate(self):
        response = self.mutate('/api/admin/settings', {'checkout_enabled': True}, 'patch')
        self.assertEqual(response.status_code, 400)
        response = server.app.test_client().post('/api/orders', json=self.order_data())
        self.assertEqual(response.status_code, 503)
        lead=server.app.test_client().post('/api/leads',json={'name':'Khách','phone':'0988123456','message':'Tư vấn','branch':'cs1','consent':True})
        self.assertEqual(lead.status_code,503)
        with server.database(True) as db:
            db.execute('DELETE FROM sims')
        self.assertEqual(self.client.get('/api/sims').json['total'], 0)

    def test_checkout_persistence_and_lookup_privacy(self):
        order, data = self.create_order()
        public = server.app.test_client()
        self.assertEqual(public.get(f'/api/sims/{self.sid}').json['status'], 'reserved')
        lookup = public.post('/api/orders/lookup', json={'id': order['id'], 'token': data['lookup_token']})
        self.assertEqual(lookup.status_code, 200)
        self.assertEqual(lookup.json['price'],1500000)
        self.assertNotIn('phone', lookup.json)
        self.assertNotIn('address', lookup.json)
        self.assertEqual(public.post('/api/orders/lookup', json={'id': order['id'], 'token': 'incorrect'}).status_code, 404)
        duplicate = public.post('/api/orders', json=data)
        self.assertEqual(duplicate.json['id'], order['id'])
        self.assertEqual(self.client.get('/api/admin/orders').json['total'], 1)
        server.initialize()
        self.assertEqual(self.client.get('/api/admin/orders').json['items'][0]['id'], order['id'])
        with server.database() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM outbox').fetchone()[0], 1)

    def test_concurrent_reservations(self):
        self.open_checkout()
        def buy(_):
            with server.app.test_client() as client:
                response = client.post('/api/orders', json=self.order_data())
                return response.status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            codes = list(pool.map(buy, range(2)))
        self.assertEqual(sorted(codes), [201, 409])
        self.assertEqual(self.client.get('/api/admin/orders').json['total'], 1)

    def test_expired_hold_and_order_transitions(self):
        order, data = self.create_order()
        with server.database(True) as db:
            db.execute('UPDATE orders SET hold_until=?', (server.now() - 1,))
        self.assertEqual(self.client.get(f'/api/sims/{self.sid}').json['status'], 'available')
        self.assertEqual(self.client.get('/api/admin/orders').json['items'][0]['status'], 'expired')
        self.assertEqual(self.mutate('/api/admin/orders/' + order['id'], {'status': 'confirmed'}, 'patch').status_code, 409)
        second = server.app.test_client().post('/api/orders', json=self.order_data()).json
        oid = second['id']
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'completed'}, 'patch').status_code, 409)
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'confirmed'}, 'patch').status_code, 200)
        self.assertEqual(self.client.get(f'/api/sims/{self.sid}').json['status'], 'sold')
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'shipping'}, 'patch').status_code, 200)
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'completed'}, 'patch').status_code, 400)
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'completed', 'paid': True}, 'patch').status_code, 200)
        self.assertEqual(self.client.get('/api/admin/dashboard').json['revenue'][0]['total'], 1500000)

    def test_cancellation_releases_stock_but_paid_requires_refund(self):
        order, _ = self.create_order()
        oid = order['id']
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'paid': True}, 'patch').status_code, 200)
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'cancelled'}, 'patch').status_code, 400)
        self.assertEqual(self.mutate('/api/admin/orders/' + oid, {'status': 'cancelled', 'paid': False}, 'patch').status_code, 200)
        self.assertEqual(self.client.get(f'/api/sims/{self.sid}').json['status'], 'available')

    def test_staff_scope_and_csrf(self):
        other_sid = self.add_sim('0987654321', 'cs2')['id']
        self.open_checkout()
        order = server.app.test_client().post('/api/orders', json=self.order_data(other_sid)).json
        staff = server.app.test_client()
        csrf = self.login(staff, 'staff1')
        self.assertEqual(staff.get('/api/admin/sims').json['total'], 1)
        self.assertEqual(staff.get('/api/admin/orders').json['total'], 0)
        self.assertEqual(self.mutate('/api/admin/orders/' + order['id'], {'status': 'confirmed'}, 'patch', staff, csrf).status_code, 403)
        self.assertEqual(self.mutate('/api/admin/settings', {'hours': 'X'}, 'patch', staff, csrf).status_code, 403)
        self.assertEqual(staff.get('/api/admin/backup').status_code, 403)
        self.assertEqual(staff.get('/api/admin/export/orders').status_code, 200)
        self.assertNotIn(order['id'], staff.get('/api/admin/export/orders').text)
        self.assertEqual(self.client.post('/api/admin/sims', json={'number':'0987654321'}).status_code, 403)
        self.assertEqual(self.client.post('/api/leads', json={}, headers={'Origin':'https://evil.test'}).status_code, 403)
        self.assertEqual(server.app.test_client().get('/api/admin/orders').status_code, 401)

    def test_staff_cannot_confirm_payment(self):
        order, _ = self.create_order()
        staff = server.app.test_client()
        csrf = self.login(staff, 'staff1')
        self.assertEqual(self.mutate('/api/admin/orders/' + order['id'], {'paid':True}, 'patch', staff, csrf).status_code, 403)
        self.assertEqual(self.mutate('/api/admin/orders/' + order['id'], {'status':'confirmed','assigned_to':2}, 'patch', staff, csrf).status_code, 200)

    def test_lead_assignment_and_schedule(self):
        self.mutate('/api/admin/settings', {'privacy_policy': 'Chính sách thử nghiệm.'}, 'patch')
        public = server.app.test_client()
        response = public.post('/api/leads', json={'name':'Khách','phone':'+84 344 639 666','message':'Tìm số đẹp','branch':'cs2','consent':True})
        self.assertEqual(response.status_code, 201)
        lid = response.json['id']
        self.assertEqual(self.mutate(f'/api/admin/leads/{lid}', {'status':'contacted','assigned_to':3,'follow_up':'2026-10-10T16:30','staff_note':'Gọi lại'}, 'patch').status_code, 200)
        item = self.client.get('/api/admin/leads').json['items'][0]
        self.assertEqual(item['phone'], '0344639666')
        self.assertEqual(item['assigned_to'], 3)
        self.assertEqual(self.mutate(f'/api/admin/leads/{lid}', {'assigned_to':2}, 'patch').status_code, 400)
        self.assertEqual(self.mutate(f'/api/admin/leads/{lid}', {'follow_up':'bad'}, 'patch').status_code, 400)

    def test_import_1000_sim_search_pagination_and_atomic_errors(self):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['number','price','type','branch'])
        for n in range(1000):
            writer.writerow(['097'+str(n).zfill(7),1000000+n,'Sim tam hoa','cs1' if n%2==0 else 'cs2'])
        response = self.mutate('/api/admin/sims/import', {'csv': output.getvalue()})
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['count'], 1000)
        data = self.client.get('/api/sims?size=24&page=2&type=Sim%20tam%20hoa').json
        self.assertEqual(data['total'],1000)
        self.assertEqual(len(data['items']),24)
        self.assertEqual(self.client.get('/api/sims?q=0970000999').json['total'],1)
        self.assertEqual(self.client.get('/api/sims?branch=cs2').json['total'],500)
        bad = 'number,price,type,branch\n0961234567,1000000,Sim đẹp,cs1\n0970000000,1000000,Sim đẹp,cs1'
        self.assertEqual(self.mutate('/api/admin/sims/import', {'csv':bad}).status_code,409)
        self.assertEqual(self.client.get('/api/sims?q=0961234567').json['total'],0)
        bad = 'number,price,type,branch\n0961234567,1000000,Sim đẹp,cs1\ninvalid,1000000,Sim đẹp,cs1'
        self.assertEqual(self.mutate('/api/admin/sims/import', {'csv':bad}).status_code,400)
        self.assertEqual(self.client.get('/api/sims').json['total'],1001)
        self.assertEqual(self.client.get('/api/sims?q=%27%20OR%201%3D1--').status_code,200)
        self.assertEqual(self.client.get('/api/sims?size=100000').status_code,400)

    def test_excel_import_export_template_and_formula_rejection(self):
        template = self.client.get('/api/admin/template.xlsx')
        self.assertEqual(template.status_code,200)
        book = load_workbook(io.BytesIO(template.data))
        sheet=book.worksheets[0]
        # Template formats 5000 rows; write explicitly to row 2.
        for index,value in enumerate(['0961234567',2500000,'Sim số đẹp','cs3'],1): sheet.cell(2,index,value)
        output=io.BytesIO();book.save(output);book.close()
        response=self.mutate('/api/admin/sims/import-xlsx',{'content':base64.b64encode(output.getvalue()).decode()})
        self.assertEqual(response.status_code,201,response.json)
        exported=self.client.get('/api/admin/export/sims.xlsx')
        self.assertEqual(exported.status_code,200)
        workbook=load_workbook(io.BytesIO(exported.data))
        values=list(workbook.active.values)
        self.assertEqual(values[1][1],'0961234567')
        workbook.close()
        formula=Workbook();formula.active.append(['number','price','type','branch']);formula.active.append(['0969999999','=1+1','Sim số đẹp','cs1']);output=io.BytesIO();formula.save(output)
        self.assertEqual(self.mutate('/api/admin/sims/import-xlsx',{'content':base64.b64encode(output.getvalue()).decode()}).status_code,400)

    def test_bulk_price_audit_and_reserved_protection(self):
        response=self.mutate('/api/admin/sims/bulk-price',{'ids':[self.sid],'price':2000000})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.client.get(f'/api/sims/{self.sid}').json['price'],2000000)
        self.create_order()
        self.assertEqual(self.mutate('/api/admin/sims/bulk-price',{'ids':[self.sid],'price':3000000}).status_code,409)
        self.assertTrue(any(r['action']=='sim.bulk_price' for r in self.client.get('/api/admin/audit').json['items']))

    def test_user_password_and_disabled_session(self):
        response=self.mutate('/api/admin/users',{'username':'newstaff','name':'Nhân viên','password':self.password,'role':'staff','branch':'cs3'})
        self.assertEqual(response.status_code,201)
        staff=server.app.test_client();csrf=self.login(staff,'newstaff')
        self.assertEqual(self.mutate('/api/admin/users/'+str(response.json['id']),{'active':False},'patch').status_code,200)
        self.assertEqual(staff.get('/api/admin/me').status_code,401)
        new=secrets.token_urlsafe(24)
        self.assertEqual(self.mutate('/api/admin/password',{'current_password':self.password,'new_password':new}).status_code,200)
        self.assertEqual(self.client.get('/api/admin/me').status_code,401)
        self.assertEqual(server.app.test_client().post('/api/login',json={'username':'owner','password':new}).status_code,200)

    def test_validation_rate_limits_and_public_file_protection(self):
        self.mutate('/api/admin/settings', {'privacy_policy': 'Chính sách thử nghiệm.'}, 'patch')
        public=server.app.test_client()
        for i in range(8):
            response=public.post('/api/leads',json={'name':'Test','phone':'0988123456','message':'Hello','branch':'cs1','consent':True})
            self.assertEqual(response.status_code,201)
        self.assertEqual(public.post('/api/leads',json={}).status_code,429)
        for path in ['/server.py','/data/store.db','/.env','/tests/test_api.py','/requirements.txt','/assets/../server.py']:
            self.assertEqual(public.get(path).status_code,404,path)
        self.assertIn("frame-ancestors 'none'",public.get('/').headers['Content-Security-Policy'])

    def test_settings_urls_policy_xss_and_seo(self):
        response=self.mutate('/api/admin/settings',{'site_url':'https://phuha.example','zalo_url':'javascript:alert(1)'},'patch')
        self.assertEqual(response.status_code,400)
        response=self.mutate('/api/admin/settings',{'site_url':'https://phuha.example','privacy_policy':'<script>alert(1)</script>','hours':'09:00–19:00'},'patch')
        self.assertEqual(response.status_code,200)
        page=self.client.get(f'/sim/{self.sid}')
        self.assertEqual(page.status_code,200)
        self.assertIn('og:title',page.text)
        self.assertIn('https://phuha.example/sim/'+str(self.sid),page.text)
        self.assertIn('Product',page.text)
        self.assertIn('/sim/'+str(self.sid),self.client.get('/sitemap.xml').text)
        self.assertEqual(self.client.get('/sim/999999').status_code,404)
        self.assertIn('Disallow: /admin',self.client.get('/robots.txt').text)

    def test_image_validation_and_reencoding(self):
        fake=base64.b64encode(b'\xff\xd8\xff'+b'not-an-image'*10).decode()
        self.assertEqual(self.mutate('/api/admin/uploads',{'content':fake}).status_code,400)
        output=io.BytesIO();Image.new('RGB',(20,20),'red').save(output,format='PNG')
        response=self.mutate('/api/admin/uploads',{'content':base64.b64encode(output.getvalue()).decode()})
        self.assertEqual(response.status_code,201,response.json)
        with self.client.get(response.json['url']) as downloaded:
            self.assertEqual(downloaded.status_code,200)
        self.assertEqual(self.mutate('/api/admin/settings',{'banner_url':response.json['url']},'patch').status_code,200)

    def test_backup_restore_consistency_and_no_sessions(self):
        order,_=self.create_order()
        backup=self.client.get('/api/admin/backup')
        self.assertEqual(backup.status_code,200)
        backup_file=Path(temporary.name)/'snapshot.sqlite3';backup_file.write_bytes(backup.data)
        with sqlite3.connect(backup_file) as db:
            self.assertEqual(db.execute('SELECT count(*) FROM sessions').fetchone()[0],0)
            self.assertEqual(db.execute('SELECT id FROM orders').fetchone()[0],order['id'])
        with server.database(True) as db:
            db.execute("UPDATE sims SET price=123 WHERE id=?",(self.sid,))
        result=subprocess.run([sys.executable,str(server.ROOT/'server.py'),'restore','--file',str(backup_file)],env=os.environ.copy(),capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        with server.database() as db:
            self.assertEqual(db.execute('SELECT price FROM sims WHERE id=?',(self.sid,)).fetchone()[0],1500000)
            self.assertEqual(db.execute('SELECT count(*) FROM sessions').fetchone()[0],0)

    def test_notifications_and_anonymous_metrics(self):
        self.create_order()
        with patch.dict(os.environ,{'SMTP_HOST':'','SMTP_FROM':''}): server.deliver_mail()
        with server.database() as db:self.assertEqual(db.execute('SELECT status FROM outbox').fetchone()[0],'pending')
        with patch.dict(os.environ,{'SMTP_HOST':'smtp.example','SMTP_FROM':'store@example.test'}),patch('smtplib.SMTP_SSL') as smtp:
            server.deliver_mail()
            smtp.return_value.send_message.assert_called_once()
        with server.database() as db:self.assertEqual(db.execute('SELECT status FROM outbox').fetchone()[0],'sent')
        self.assertEqual(self.client.post('/api/view',json={'page':'/kho-sim'}).status_code,200)
        self.assertEqual(self.client.get('/api/admin/dashboard').json['views'][0]['count'],1)


if __name__=='__main__':
    unittest.main(verbosity=2)
